#!/usr/bin/env python3
"""Tests for the per-subject build. Run: python3 tests/build_subjects_test.py

Checks: manifest and per-subject files, id scheme, registry numbers stay stable when a subject is
added, a conflict inside a subject refuses, a cross-subject duplicate builds and is reported,
stale subject data is removed.
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tests"))
import make_sample_subjects as sample  # noqa: E402

BUILD = [sys.executable, str(ROOT / "tools" / "build_data.py")]
failures = []


def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        failures.append(msg)


def build(src, out, reg, report, *extra):
    return subprocess.run(BUILD + ["--source", str(src), "--out-dir", str(out), "--registry", str(reg),
                                   "--report", str(report), *extra], capture_output=True, text=True)


def load(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    src, out, reg, rep = tmp / "src", tmp / "out", tmp / "subjects.json", tmp / "report.txt"
    sample.main(str(src))

    # 1. basic build (near threshold lowered so the river pair in geography is reported)
    r = build(src, out, reg, rep, "--near", "0.7")
    check(r.returncode == 0, "build succeeds on two subjects")
    manifest = load(out / "manifest.json")
    names = [s["name"] for s in manifest["subjects"]]
    check(names == ["جغرافيا", "علوم"], f"manifest lists both subjects in file order: {names}")
    check([s["count"] for s in manifest["subjects"]] == [12, 12], "each subject has 12 questions")
    geo = load(out / manifest["subjects"][0]["file"])
    sci = load(out / manifest["subjects"][1]["file"])
    geo_ids = [i["id"] for i in geo["items"]]
    sci_ids = [i["id"] for i in sci["items"]]
    check(geo_ids[0] == 1 * 100000 + 1 and sci_ids[0] == 2 * 100000 + 1, "ids are number*100000+sequence")
    check(len(set(geo_ids) | set(sci_ids)) == 24, "all 24 ids are globally unique")
    reg_data = load(reg)
    check([(x["file"], x["number"]) for x in reg_data] == [("جغرافيا.xlsx", 1), ("علوم.xlsx", 2)], "registry assigned numbers 1 and 2")

    # 2. report content: cross-subject duplicate is info, near-duplicate within a subject is reported
    text = rep.read_text(encoding="utf-8")
    check("same question in subjects جغرافيا, علوم (DIFFERENT answers)" in text, "cross-subject duplicate reported as info")
    check("near-duplicate in 'جغرافيا'" in text, "near-duplicate inside geography reported")

    # 3. adding a subject keeps existing ids and numbers
    before = {i["id"]: i["q"] for i in geo["items"] + sci["items"]}
    sample.main(str(src), third=True)
    r = build(src, out, reg, rep)
    check(r.returncode == 0, "build succeeds with a third subject")
    m2 = load(out / "manifest.json")
    check([s["number"] for s in m2["subjects"]] == [1, 2, 3], f"numbers stay 1,2 and new one is 3: {[s['number'] for s in m2['subjects']]}")
    after = {}
    for s in m2["subjects"]:
        for i in load(out / s["file"])["items"]:
            after[i["id"]] = i["q"]
    check(all(after.get(k) == v for k, v in before.items()), "existing question ids unchanged after adding a subject")
    v1 = {s["number"]: s["version"] for s in manifest["subjects"]}
    v2 = {s["number"]: s["version"] for s in m2["subjects"]}
    check(v1[1] == v2[1] and v1[2] == v2[2], "unchanged subjects keep their content version")

    # 4. removing a file: its data is removed, its number stays reserved
    (src / "تاريخ.xlsx").unlink()
    r = build(src, out, reg, rep)
    check(not (out / "3.json").exists(), "stale subject data removed")
    check(any(x["number"] == 3 for x in load(reg)), "removed subject's number stays reserved in registry")
    sample.main(str(src), third=True)
    build(src, out, reg, rep)
    check(load(out / "manifest.json")["subjects"][-1]["number"] == 3, "re-adding the file gets its old number back")

    # 5. conflict inside one subject refuses and writes nothing new
    wb = openpyxl.load_workbook(src / "علوم.xlsx")
    wb.active.append(["ما هو الكوكب الأقرب إلى الشمس؟", "الزهرة"])
    wb.save(src / "علوم.xlsx")
    out2 = tmp / "out2"
    r = build(src, out2, tmp / "reg2.json", rep)
    check(r.returncode != 0 and "REFUSED" in r.stderr + r.stdout, "conflict inside a subject refuses to build")
    check(not (out2 / "manifest.json").exists(), "nothing is written when the build is refused")
    r = build(src, out2, tmp / "reg2.json", rep, "--allow-conflicts")
    check(r.returncode == 0 and (out2 / "manifest.json").exists(), "--allow-conflicts builds anyway")

    # 6. several sheets in one workbook merge into one subject
    wb = openpyxl.load_workbook(src / "جغرافيا.xlsx")
    ws2 = wb.create_sheet("إضافي")
    ws2.append(["السؤال", "الإجابة"])
    ws2.append(["ما هي عاصمة اليابان؟", "طوكيو"])
    wb.save(src / "جغرافيا.xlsx")
    (src / "علوم.xlsx").unlink()
    sample.main(str(src))  # rewrites both, so re-add the extra sheet to geography
    wb = openpyxl.load_workbook(src / "جغرافيا.xlsx")
    ws2 = wb.create_sheet("إضافي")
    ws2.append(["السؤال", "الإجابة"])
    ws2.append(["ما هي عاصمة اليابان؟", "طوكيو"])
    wb.save(src / "جغرافيا.xlsx")
    out3 = tmp / "out3"
    r = build(src, out3, tmp / "reg3.json", rep)
    geo_entry = next(s for s in load(out3 / "manifest.json")["subjects"] if s["name"] == "جغرافيا")
    check(r.returncode == 0 and geo_entry["count"] == 13, f"two sheets merge into one subject (13 questions): {geo_entry['count']}")

print()
print(f"{len(failures)} failure(s)" if failures else "ALL PASS")
sys.exit(1 if failures else 0)
